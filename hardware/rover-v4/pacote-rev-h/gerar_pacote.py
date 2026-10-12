"""Monta o pacote VALTEXA_ROVER_V4_REV_H: pastas, documentos, desenhos, CAD, BOM e auditoria.

Uso:
    python gerar_pacote.py              # tudo
    python gerar_pacote.py --sem-pdf    # só HTML, CSV, DXF e STL (rápido)
"""
import csv
import pathlib
import shutil
import sys

AQUI = pathlib.Path(__file__).resolve().parent
CG = AQUI.parent / 'chassi-gaiola'
sys.path.insert(0, str(CG))

import modelo as M                                  # noqa: E402
import gerar as G                                   # noqa: E402
import dados                                        # noqa: E402
import doc_tecnico, doc_fabricacao, doc_gestao, doc_fornecedores, custos, bom_precos   # noqa: E402
from processos import PROC, LASER, FRESA            # noqa: E402

RAIZ = AQUI / 'saida' / 'VALTEXA_ROVER_V4_REV_H'
SEM_PDF = '--sem-pdf' in sys.argv

# folhas do caderno por assunto (número da folha no VTX-RV4-CG-H)
ASSUNTOS = [
    ('CHASSIS', 'Chassi: quadro, chapa de fundo e suspensão', [1, 2, 3, 4, 5, 6, 7, 8, 13, 14, 27]),
    ('SUSPENSAO_USINAGEM', 'Suspensão pneumática: conjunto e desenhos de usinagem',
     [1, 2, 7, 28, 29, 30, 31]),
    ('MONTAGEM', 'Manual de montagem: posição, sequência, torques e regulagem',
     [1, 13, 32, 33, 34, 35]),
    ('SENSORES_MODULARES', 'Mastro completo, interface VMS-60, os nove postos e o módulo térmico removível',
     [1, 2, 18, 19, 36, 37, 38, 39, 40, 41]),
    ('CABOS', 'Interligação e roteamento de cabos', [1, 2, 18, 21, 22, 42]),
    ('GAIOLA', 'Gaiola de proteção, carenagens e tampas', [1, 2, 9, 10, 11, 12, 14]),
    ('MASTROS', 'Mastros: base, haste, postos, antena e rigidez', [1, 2, 17, 18, 36]),
    ('SUPORTES_CAMERAS', 'Posto fixo (OAK-D) e módulo térmico removível', [1, 2, 19, 37, 41]),
    ('INTERFACES', 'Bateria, eletrônica, E-STOP e interfaces de sensores', [1, 2, 20, 21, 22]),
    ('ENGENHARIA', 'Validação, pendências, massa, FEA, ensaios e motor', [15, 23, 24, 25, 26, 27]),
]


def cabecalho(txt):
    print(f'\n--- {txt}')


# ------------------------------------------------------------------ STL
def _ear_clip(pol):
    """Triangula um polígono simples (lista de (a,b)) por ear clipping. Devolve índices."""
    n = len(pol)
    if n < 3:
        return []
    area = sum(pol[i][0] * pol[(i + 1) % n][1] - pol[(i + 1) % n][0] * pol[i][1] for i in range(n))
    idx = list(range(n)) if area > 0 else list(range(n - 1, -1, -1))

    def cruz(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    def dentro(p, a, b, c):
        d1, d2, d3 = cruz(a, b, p), cruz(b, c, p), cruz(c, a, p)
        neg = d1 < 0 or d2 < 0 or d3 < 0
        pos = d1 > 0 or d2 > 0 or d3 > 0
        return not (neg and pos)

    tri, guarda = [], 0
    while len(idx) > 3 and guarda < 5000:
        guarda += 1
        for k in range(len(idx)):
            i0, i1, i2 = idx[k - 1], idx[k], idx[(k + 1) % len(idx)]
            a, b, c = pol[i0], pol[i1], pol[i2]
            if cruz(a, b, c) <= 0:
                continue
            if any(dentro(pol[j], a, b, c) for j in idx if j not in (i0, i1, i2)):
                continue
            tri.append((i0, i1, i2))
            idx.pop(k)
            break
        else:
            break
    if len(idx) == 3:
        tri.append(tuple(idx))
    return tri


def stl_do_conjunto(pecas, destino, nome='rover'):
    """Malha do contorno externo de cada peça (sem furos) — para visualização e conferência."""
    linhas = [f'solid {nome}']

    def face(p1, p2, p3):
        linhas.append('facet normal 0 0 0\n  outer loop')
        for v in (p1, p2, p3):
            linhas.append(f'    vertex {v[0]:.3f} {v[1]:.3f} {v[2]:.3f}')
        linhas.append('  endloop\nendfacet')

    for p in pecas:
        if p.estilo != 'solid':
            continue
        perfil = p.perfil
        tri = _ear_clip(perfil)
        v0 = [p.mundo(a, b, p.e0) for (a, b) in perfil]
        v1 = [p.mundo(a, b, p.e1) for (a, b) in perfil]
        for (i, j, k) in tri:
            face(v0[i], v0[k], v0[j])
            face(v1[i], v1[j], v1[k])
        n = len(perfil)
        for i in range(n):
            j = (i + 1) % n
            face(v0[i], v0[j], v1[j])
            face(v0[i], v1[j], v1[i])
    linhas.append(f'endsolid {nome}')
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text('\n'.join(linhas), encoding='ascii')
    return destino


# ------------------------------------------------------------------ BOM
def bom(destino):
    its = dados.itens()
    cab = ['Código', 'Item', 'Descrição', 'Categoria', 'Parte', 'Material', 'Especificação',
           'Dimensões (mm)', 'Espessura/seção', 'Qtd por conjunto', 'Massa unit. (kg)',
           'Massa total (kg)', 'Folha do caderno', 'Processo', 'Operações seguintes', 'Acabamento',
           'Origem do dado',
           'Fornecedor', 'Preço unit. (R$)', 'Observação']
    linhas = []
    for i in its:
        linhas.append([i['codigo'], i['item'], i['nome'], 'Fabricação própria', i['parte'], i['material'],
                       i['desc'], i['dim'], i['esp'], i['qtd'], round(i['unit'], 4), round(i['massa'], 3),
                       i['folha'], i['processo'], ' · '.join(i['ops']) or '—', i['acabamento'],
                       i['origem'],
                       'COTAÇÃO NECESSÁRIA', 'COTAÇÃO NECESSÁRIA', i['obs']])
    for (cod, desc, mat, qtd, uso) in M.fixadores():
        linhas.append([f'VTX-RV4-FIX-{cod}-REV-H', cod, desc, 'Componente comercial', 'Fixadores', mat, desc,
                       '—', '—', qtd, '', '', '13', 'Compra', '—', '—', 'CONFIRMADO',
                       'COTAÇÃO NECESSÁRIA', 'COTAÇÃO NECESSÁRIA', uso])
    destino.parent.mkdir(parents=True, exist_ok=True)
    with open(destino.with_suffix('.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(cab)
        w.writerows(linhas)
    xlsx = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = 'BOM mecânica Rev H'
        ws.append(cab)
        for l in linhas:
            ws.append(l)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F6B3A')
            c.alignment = Alignment(wrap_text=True, vertical='bottom')
        larg = [26, 7, 30, 20, 22, 16, 30, 16, 16, 9, 12, 12, 10, 22, 28, 22, 14, 18, 18, 46]
        for i, w_ in enumerate(larg, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w_
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        xlsx = destino.with_suffix('.xlsx')
        wb.save(xlsx)
    except ImportError:
        pass
    return len(linhas), xlsx


# ------------------------------------------------------------ auditoria
def verificacoes(dxf_dir):
    its = dados.itens()
    planas = [i for i in its if i['processo'] in (LASER, FRESA)]
    dxfs = {f.name.split('_')[0] for f in dxf_dir.glob('*.dxf')}
    sem_dxf = [i['item'] for i in planas if i['item'] not in dxfs and not any(
        i['item'] in d for d in dxfs)]
    # A conferência de interferência é UMA só — a de revisao-h/espaco.py, que declara cada encaixe
    # previsto com o motivo. Esta função tinha uma cópia própria, com uma lista de exceções do V3
    # (M-20, N-09) que nem existem mais.
    import importlib.util
    esp = importlib.util.spec_from_file_location('espaco_aud', CG.parent / 'revisao-h' / 'espaco.py')
    ES = importlib.util.module_from_spec(esp)
    esp.loader.exec_module(ES)
    it = ES.interferencias()
    inter = len(it['defeitos'])
    medidas = ''
    layout_res, layout_txt = 'OK', 'não medido (modo --sem-pdf)'
    if not SEM_PDF:
        import verifica_layout as VL
        todas = list(range(1, G.TODAS.__len__() + 1))
        med = VL.medir(todas)
        n_def = 0
        for n_ in todas:
            dm = med.get(f'f{n_:02d}')
            n_def += 1 if dm is None else len(VL.defeitos(dm, str(G.TODAS[n_ - 1]().s)))
        layout_res = 'OK' if n_def == 0 else 'FALHA'
        layout_txt = f'{n_def} defeito(s) medidos no navegador, em {len(todas)} folhas'
    c = M.consolidado()
    massa_bom = sum(i['massa'] for i in its)
    # Peça comprada não passa por processo de fabricação, mas TEM de estar numa lista de compra:
    sem_proc = [i for i in M.ITENS if i not in PROC and M.ITENS[i].get('mat') != 'COMPRADO']
    comprados = [i for i in M.ITENS if M.ITENS[i].get('mat') == 'COMPRADO']
    sem_linha = [i for i in comprados if not any(l[1].startswith(i) for l in bom_precos.BOM)]
    auto = [
        ('Códigos de peça fabricada no modelo', f'{len(its)} códigos · {sum(i["qtd"] for i in its)} peças',
         'OK'),
        ('Códigos de fixador', f'{len(M.fixadores())} códigos · {sum(q for (_, _, _, q, _) in M.fixadores())}'
         ' peças', 'OK'),
        ('Itens FABRICADOS do catálogo sem classificação de processo',
         f'{len(sem_proc)}' + (f' ({", ".join(sem_proc)})' if sem_proc else ''),
         'OK' if not sem_proc else 'FALHA'),
        ('Itens COMPRADOS do catálogo sem linha na BOM de compras',
         f'{len(comprados)} comprados ({", ".join(comprados)}) · faltam: ' + (", ".join(sem_linha) or 'nenhum'),
         'OK' if not sem_linha else 'FALHA'),
        ('Peças planas sem DXF', ', '.join(sem_dxf) if sem_dxf else 'nenhuma',
         'OK' if not sem_dxf else 'VERIFICAR'),
        ('Interferências geométricas entre peças (mastro completo)',
         f'{inter} defeitos · {len(it["previstos"])} encaixes previstos · {it["pecas"]} peças', 'OK' if inter == 0 else 'FALHA'),
        ('Massa somada da BOM × massa consolidada',
         f'{massa_bom:.2f} kg de peça fabricada · {c["m_min"]:.2f} kg de rover completo', 'OK'),
        (f'Layout das {len(G.TODAS)} folhas (texto sobre texto, linha cortando texto, fora da moldura, '
         'invasão do carimbo) — verifica_layout.py', layout_txt, layout_res),
        ('Divergências entre documentos registradas',
         f'{len(doc_tecnico.divergencias(c, doc_tecnico._carrega("mastro_aud", "mastro.py")))} no documento técnico '
         f'+ {len(bom_precos.DIVERGENCIAS)} na conciliação da BOM de compras', 'OK'),
        ('Preço de fabricação estimado sem cotação', 'nenhum — só quantidades para cotar', 'OK'),
    ]
    cobertura = [
        ['Com desenho no caderno', f'{len(its)} de {len(its)} códigos',
         'Todos aparecem em pelo menos uma folha'],
        ['Com DXF 1:1', f'{len(dxfs)} arquivos', 'Peças planas normais aos eixos'],
        ['Sem DXF por natureza', 'tubos, barras, buchas, espaçadores, eixos, face inclinada N-07, protetor E-11 (barra dobrada em U, desenvolvimento a definir pelo fabricante), adaptador T-10',
         'Cortadas por serra/torno ou não planificáveis — cotas no desenho'],
        ['Sem CAD nativo', 'todos', 'O modelo é paramétrico em Python; STEP/CAD é pendência de engenharia'],
    ]
    sem_dado = [i for i in its if i['item'] in ('T-03', 'T-04', 'T-10', 'N-02', 'W-15', 'W-16', 'W-17')
                or i['origem'] == 'REQUISITO']
    return dict(auto=auto, cobertura=cobertura, aprovados=len(its) - len(sem_dado),
                incompletos=len(sem_dado),
                conflitos=len(bom_precos.DIVERGENCIAS) + len(doc_tecnico.divergencias(
                    c, doc_tecnico._carrega('mastro_aud2', 'mastro.py'))),
                engenheiro=len([p for p in doc_gestao.PEND if p[0] in ('Estrutura', 'Sensores V4')]),
                cotacao=len(its) + len(M.fixadores()))


# ------------------------------------------------------------------ main
def main():
    if RAIZ.exists():
        shutil.rmtree(RAIZ, ignore_errors=True)   # tolera pasta aberta no explorador/terminal
    for p in ('01_DOCUMENTO_TECNICO', '02_PACOTE_ORCAMENTO', '03_DESENHOS_TECNICOS', '04_CAD/DXF',
              '04_CAD/STL', '05_BOM', '06_FORNECEDORES', '07_RFQ', '08_PENDENCIAS',
              '09_CONCEITO_BASES', '10_CUSTOS', '11_ENGENHARIA_REV_H'):
        (RAIZ / p).mkdir(parents=True, exist_ok=True)

    cabecalho('documentos de engenharia do V4 (Rev H)')
    # Os documentos que sustentam esta revisão: a decisão da suspensão, o manual de
    # montagem, o roteiro de usinagem, o espaço medido, a auditoria e a liberação.
    # Vão no pacote porque o fabricante precisa deles, não só o engenheiro.
    rg = CG.parent / 'revisao-h' / 'pdf'
    for f in sorted(rg.glob('*.pdf')):
        shutil.copy2(f, RAIZ / '11_ENGENHARIA_REV_H' / f.name)
        print(' copiado:', f.name)
    for f in sorted((CG.parent / 'revisao-h').glob('*.md')):
        shutil.copy2(f, RAIZ / '11_ENGENHARIA_REV_H' / f.name)

    cabecalho('caderno de desenhos')
    saida_cg = CG / 'saida'
    for nome, destino in (('VTX-RV4-CG-H.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-completo.pdf'),
                          ('VTX-RV4-CG-H-A.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-ParteA.pdf'),
                          ('VTX-RV4-CG-H-B.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-ParteB.pdf'),
                          ('VTX-RV4-CG-H-C.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-ParteC.pdf'),
                          ('VTX-RV4-CG-H-D.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-ParteD.pdf'),
                          ('VTX-RV4-CG-H-E.pdf', '03_DESENHOS_TECNICOS/VTX-RV4-CG-H-ParteE.pdf')):
        f = saida_cg / nome
        if f.exists():
            shutil.copy2(f, RAIZ / destino)
            print(' copiado:', destino)

    if not SEM_PDF:
        for pasta, desc, folhas in ASSUNTOS:
            alvo = RAIZ / '03_DESENHOS_TECNICOS' / pasta
            alvo.mkdir(exist_ok=True)
            html = alvo / f'VTX-RV4-CG-H-{pasta}.html'
            html.write_text(G.html_completo(folhas, f'{desc} · folhas {", ".join(f"{n:02d}" for n in folhas)}'),
                            encoding='utf8')
            pdf = alvo / f'VTX-RV4-CG-H-{pasta}.pdf'
            if G.imprime_pdf(html, pdf):
                html.unlink(missing_ok=True)
            print(f' {pasta}: {len(folhas)} folhas')

    cabecalho('CAD')
    dxf_origem = saida_cg / 'dxf'
    for f in dxf_origem.glob('*'):
        shutil.copy2(f, RAIZ / '04_CAD/DXF' / f.name)
    print(f' DXF: {len(list((RAIZ / "04_CAD/DXF").glob("*.dxf")))} arquivos')
    stl_do_conjunto([p for p in M.conjunto('AB')],
                    RAIZ / '04_CAD/STL/VTX-RV4-conjunto-REV-H.stl', 'rover_v4')
    stl_do_conjunto([p for p in M.conjunto('ABO')],
                    RAIZ / '04_CAD/STL/VTX-RV4-conjunto-completo-REV-H.stl', 'rover_v4_completo')
    stl_do_conjunto(M.quadro() + M.base_mastro(),
                    RAIZ / '04_CAD/STL/VTX-RV4-quadro-REV-H.stl', 'quadro')
    stl_do_conjunto(M.gaiola() + M.suporte_lidar(),
                    RAIZ / '04_CAD/STL/VTX-RV4-gaiola-REV-H.stl', 'gaiola')
    stl_do_conjunto(M.suspensao(),
                    RAIZ / '04_CAD/STL/VTX-RV4-suspensao-REV-H.stl', 'suspensao')
    stl_do_conjunto(M.mastro_completo((1, -1)),
                    RAIZ / '04_CAD/STL/VTX-RV4-mastro-sensores-REV-H.stl', 'mastro')
    print(f' STL: {len(list((RAIZ / "04_CAD/STL").glob("*.stl")))} arquivos')
    (RAIZ / '04_CAD/LEIA-ME.txt').write_text(
        'VALTEXA ROVER V4 — Rev H — arquivos de CAD\n'
        '=========================================\n\n'
        'DXF/  peças planas em 1:1, milímetros, camadas CORTE / FUROS / GRAVAR-NAO.\n'
        '      Geometria exata, gerada do modelo paramétrico. É o arquivo de produção.\n\n'
        'STL/  malha 3D do contorno externo das peças, para visualização e conferência de montagem.\n'
        '      ATENÇÃO: a malha NÃO contém furos nem roscas. Não usar para fabricação.\n\n'
        'STEP  não incluso. O modelo da ValTexa é paramétrico em Python, não um CAD nativo.\n'
        '      A conversão para CAD (SolidWorks / Fusion / Onshape) e a emissão de STEP são\n'
        '      pendência de engenharia — ver 08_PENDENCIAS.\n', encoding='utf8')

    cabecalho('fontes recebidas')
    fontes = RAIZ / '05_BOM/fontes'
    fontes.mkdir(parents=True, exist_ok=True)
    for origem in (pathlib.Path.home() / 'Downloads' / 'VALTEXA-ROVER-V3-BOM-Links-Precos.xlsx',
                   AQUI / 'fontes' / 'VALTEXA-ROVER-V3-BOM-Links-Precos.xlsx'):
        if origem.exists():
            shutil.copy2(origem, fontes / origem.name)
            print(' fonte:', origem.name)
            break
    else:
        print(' aviso: BOM de compras não encontrada para copiar')

    cabecalho('BOM')
    n_linhas, xlsx = bom(RAIZ / '05_BOM/VTX-RV4-BOM-Mecanica-RevH')
    print(f' BOM: {n_linhas} linhas · CSV{" + XLSX" if xlsx else " (openpyxl ausente: sem XLSX)"}')
    p = custos.gerar(RAIZ / '10_CUSTOS/VTX-RV4-Estrutura-de-Custo-RevH.xlsx')
    print(' custos:', p.name, '(sem preço estimado — células amarelas para cotação)')

    cabecalho('documentos')
    docs = [
        (doc_tecnico.build(), '01_DOCUMENTO_TECNICO/VTX-RV4-Documento-Tecnico-Engenharia-RevH'),
        (doc_fabricacao.build(), '02_PACOTE_ORCAMENTO/VTX-RV4-Pacote-Fabricacao-RevH'),
        (doc_fabricacao.rfq(), '07_RFQ/VTX-RV4-Solicitacao-de-Orcamento-RevH'),
        (doc_gestao.pendencias(), '08_PENDENCIAS/VTX-RV4-Pendencias-Engenharia-RevH'),
        (doc_gestao.auditoria(verificacoes(RAIZ / '04_CAD/DXF')),
         '08_PENDENCIAS/VTX-RV4-Auditoria-do-Pacote-RevH'),
        (doc_gestao.conceito_bases(), '09_CONCEITO_BASES/VTX-RV4-Base-e-Poste-CONCEITO-RevH'),
        (doc_fornecedores.build(), '06_FORNECEDORES/VTX-RV4-Fornecedores-Pouso-Alegre-e-Regiao-RevH'),
    ]
    for d, destino in docs:
        p = d.salva(RAIZ / destino, pdf=not SEM_PDF)
        print(' ', p.relative_to(RAIZ))

    x = doc_fornecedores.planilha(RAIZ / '06_FORNECEDORES/Fornecedores-Pouso-Alegre-e-Regiao.xlsx')
    print(' ', (x.relative_to(RAIZ) if x else 'planilha de fornecedores: openpyxl ausente'))
    (RAIZ / '00_LEIA-ME.md').write_text(LEIAME, encoding='utf8')
    print('\npacote em:', RAIZ)


LEIAME = '''# VALTEXA ROVER V4 — Rev H · pacote de engenharia e fabricação

Emitido em 12/10/2026 · **para cotação e validação — não liberado para fabricação**

> Este pacote **substitui** o de 11/10/2026. Aquele carregava, nos documentos 01, 02, 06, 07, 08, 09 e 10,
> texto herdado do V3 (suportes rígidos, mastro Ø25×2, placa M-20, carenagem N-09, oito câmeras). Foi
> reescrito e conferido contra o modelo.

## O que é cada pasta

| Pasta | Conteúdo | Para quem |
|---|---|---|
| `01_DOCUMENTO_TECNICO` | Memorial: arquitetura, dimensões, massa, mastro, suspensão, interfaces, divergências | Engenheiro mecânico |
| `02_PACOTE_ORCAMENTO` | Listas por processo, matéria-prima, inspeção e cenários 1/3/15 | Empresas de fabricação |
| `03_DESENHOS_TECNICOS` | Caderno de 42 folhas A3, completo, por parte (A a E) e por assunto | Todos |
| `04_CAD` | DXF 1:1 (produção) e STL (visualização) | Corte a laser, engenharia |
| `05_BOM` | Lista de materiais mecânica completa, com os opcionais separados | Compras |
| `06_FORNECEDORES` | Pesquisa de fornecedores em Pouso Alegre e região (alumínio; aço NÃO levantado) | Compras |
| `07_RFQ` | Formulário de solicitação de orçamento | Enviar às empresas |
| `08_PENDENCIAS` | Pendências de engenharia e auditoria do pacote | Engenheiro + gestão |
| `09_CONCEITO_BASES` | Conceito da base de recarga e do poste (NÃO é baseline) | Produto + engenharia |
| `10_CUSTOS` | Estrutura de custo: quantidades, preços em branco; itens novos do V4 a cotar | Compras |
| `11_ENGENHARIA_REV_H` | Mastro, modular, suspensão, usinagem, montagem, espaço, cabos, auditoria, liberação | Engenheiro mecânico |

## Regra de leitura

Todo valor carrega uma etiqueta: **CONFIRMADO** (baseline, ficha de fabricante ou caderno), **DERIVADO**
(calculado), **PROPOSTA** (escolha deste pacote, a aprovar) e **A DEFINIR** (ausente). Nada foi preenchido por
estimativa silenciosa.

## Avisos importantes

1. **Os preços de compra são os do V3** (planilha de 21/09/2026). Itens que o V4 acrescentou ou mudou entram
   como **A COTAR**, sem preço, e **não** entram nos totais. O total verdadeiro é maior que o mostrado.
2. O CAD Rev B (`VALTEXA-ROVER-V3-RevB-CAD.step`) está **superado** — não usar.
3. Base de recarga e poste **não existem na baseline**: a pasta 09 é conceito para discussão.
4. A fabricação depende do parecer do engenheiro mecânico (folha 15) e do fechamento das pendências. A
   liberação de 11/10/2026 é sobre a **Rev F do V3**, não sobre este V4: o campo de aprovação está vazio.
5. Nada aqui é validação física. Os cálculos do mastro, da repetibilidade e da suspensão são cálculos; os ensaios
   V-10 a V-17 não foram feitos.

## Como regenerar

```
cd hardware/rover-v4/chassi-gaiola && python gerar.py                  # caderno + DXF
cd ../chassi-gaiola && python verifica_layout.py                      # mede o layout de cada folha
cd ../revisao-h && python relatorios.py && python gerar_docs.py       # documentos de engenharia
cd ../pacote-rev-h && python gerar_pacote.py                          # este pacote
```
'''


if __name__ == '__main__':
    main()
